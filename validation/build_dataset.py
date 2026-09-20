#!/usr/bin/env python3
"""Build the complete SciML-MD validation DATA workbook from every result JSON."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

V = Path("/home/claude/SciML-MD/validation")
FONT = "Arial"
INK = "1A1A1A"
HDR = PatternFill("solid", fgColor="1F3B57")
SUB = PatternFill("solid", fgColor="E8EEF4")
BAND = PatternFill("solid", fgColor="F7F9FB")
GOOD = PatternFill("solid", fgColor="E3F2E8")
BAD = PatternFill("solid", fgColor="FBE6DC")
WARN = PatternFill("solid", fgColor="FDF3E0")
thin = Side(style="thin", color="C9D3DC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
B = Font(name=FONT, size=10, bold=True, color=INK)
N = Font(name=FONT, size=10, color=INK)
SM = Font(name=FONT, size=9, color="5A5A5A")
TITLE = Font(name=FONT, size=13, bold=True, color=INK)

wb = Workbook()
first = True


def load(name):
    p = V / name
    return json.loads(p.read_text()) if p.exists() else None


def sheet(name, title, subtitle, headers, widths):
    global first
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    first = False
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SM
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = H; c.fill = HDR; c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    return ws


def rows(ws, data, start=4, wrap=(), num_fmt=None):
    for r, row in enumerate(data, start):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = N; cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in wrap))
            if (r - start) % 2 == 1:
                cell.fill = BAND
            if num_fmt and isinstance(val, float):
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="center", vertical="top")
    return start + len(data)


def flag(ws, col, r0, r1, good_vals=("PASS",), bad_vals=("FAIL", "OPEN")):
    for r in range(r0, r1):
        c = ws.cell(row=r, column=col)
        if c.value in good_vals:
            c.fill = GOOD; c.font = B
        elif c.value in bad_vals:
            c.fill = BAD; c.font = B


# =====================================================================
# 0. README
# =====================================================================
ws = sheet("00_README", "SciML-MD — Complete Validation Dataset",
           "Every numerical result produced by the Aug 2026 validation campaign. "
           "Branch: validation-fixes. Each sheet names its source file so any value is traceable.",
           ["Sheet", "Contents", "Source file", "n"], [22, 62, 40, 8])
readme = [
    ["01_Study_Design", "Pre-registered criteria for every phase, and their outcomes", "(this document)", ""],
    ["02_Defects", "All defects found, root cause, measured impact, fix status", "VALIDATION_REPORT.md", "12"],
    ["03_L1_Numerical", "Active-set mismatch, tie spread, residue-vs-RMSF, unit checks", "numerical_results_BASELINE.json", ""],
    ["04_L1_EdgeCases", "25 edge cases, baseline vs post-fix behaviour", "edge_results_BASELINE/_POSTFIX.json", "25"],
    ["05_Accuracy", "Rare-case AUROC/AUPRC per protein per channel, before/after", "rare_case_results_BASELINE/_POSTFIX.json", ""],
    ["06_L2_MSM_Validity", "Chapman-Kolmogorov, implied timescales, bootstrap CIs, VAMP-2", "layer2_results.json", "4 systems"],
    ["07_L3_External", "B-factor correlation and baseline-method comparison", "layer3_results.json", ""],
    ["08_Revalidation", "Adversarial re-check of our own conclusions (errata E-1..E-6)", "revalidation_results.json", "8 checks"],
    ["09_Phase1_Seeded", "Seeded thermodynamic anomaly (450K frames in 300K background)", "phase1_results.json", "3 modes"],
    ["10_Phase2_Scramble", "Temporal scrambling - the decisive kinetic test", "phase2_results.json", "3 modes"],
    ["11_Phase3_ATLAS", "Scale-up pilot on ATLAS long trajectories", "phase3_results.json", "1 protein"],
    ["12_D09_Impact", "Effect of the smoothing fix on earlier accuracy results", "d09_impact.json", "4 systems"],
    ["13_Limitations", "Every known limitation and open question, stated explicitly", "(this document)", ""],
]
end = rows(ws, readme, wrap=(2,))
ws.cell(row=end + 1, column=1, value="PROVENANCE").font = B
prov = [
    ["Pipeline", "tICA -> KMeans -> reversible MLE MSM -> multi-signal anomaly scoring (rarity 1-pi, transition surprise -log P, kNN density in tICA space), median fusion of rank-normalised channels"],
    ["Environment", "Python 3.11.15, deeptime 0.4.5, mdtraj 1.11.1, scikit-learn 1.8.0, OpenMM 8.5.2, PDBFixer 1.12.0"],
    ["Reproduce", "All scripts in validation/ are re-runnable; see 01_Study_Design for the command per phase"],
    ["Key caveat", "See 13_Limitations before drawing any conclusion. Several headline claims were RETRACTED by our own re-validation (sheet 08)."],
]
r = end + 2
for a, b in prov:
    ws.cell(row=r, column=1, value=a).font = B
    c = ws.cell(row=r, column=2, value=b); c.font = N
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

# =====================================================================
# 1. STUDY DESIGN
# =====================================================================
ws = sheet("01_Study_Design", "Pre-registered criteria and outcomes",
           "Criteria were fixed BEFORE each run and are reported as-is. None were adjusted after seeing results.",
           ["Phase", "What it tests", "Ground truth source", "Criterion", "Threshold",
            "Observed", "Outcome", "Command"], [10, 34, 34, 34, 16, 14, 10, 46])
p1 = load("phase1_results.json") or {}
p2 = load("phase2_results.json") or {}
p3 = load("phase3_results.json") or {}
a1 = p1.get("acceptance", {})
a2 = p2.get("acceptance", {})
a3 = p3.get("acceptance", {})
design = [
    ["Layer 1", "Implementation correctness of the three signals",
     "Hand-computed values on an analytic 3-state Markov chain", "Formulas match to machine precision",
     "exact", "exact", "PASS", "python validation/baseline_numerical.py"],
    ["Layer 1", "Behaviour on degenerate inputs", "Constructed edge cases",
     "No silent wrong output; fail-fast where undefined", "qualitative", "23 ok / 2 fail-fast", "PASS",
     "python validation/baseline_edge_cases.py"],
    ["Accuracy", "Rare-state / rare-transition detection", "Operational: bottom-quartile-pi states; transitions seen <=1x",
     "Post-fix AUROC exceeds baseline", "improvement", "+0.10 to +0.35", "PASS",
     "python validation/baseline_rare_case.py"],
    ["Layer 2", "MSM statistical validity", "Standard MSM battery (Prinz et al. 2011)",
     "CK within CI; ITS plateau; pi CI < 0.5x value; VAMP-2 below dim", "see sheet 06",
     "sampling-limited", "FAIL", "python validation/layer2_msm_validity.py"],
    ["Layer 3", "External validity vs experiment", "Crystallographic B-factors (RCSB)",
     "Correlation with per-residue score", "rho >= 0.5", "0.595 / 0.573", "RETRACTED",
     "python validation/layer3_external.py"],
    ["Revalid.", "Adversarial re-check of our own claims", "Re-derivation from first principles",
     "Claims survive falsification attempt", "qualitative", "3 of 6 retracted", "PARTIAL",
     "python validation/revalidation.py"],
    ["Phase 1", "Seeded THERMODYNAMIC anomaly", "450K frames spliced into 300K background (splice list)",
     "A1 fused burst AUROC >= 0.80", "0.80", a1.get("A1_fused_burst_auroc_ge_0.80", {}).get("value"),
     "PASS" if a1.get("A1_fused_burst_auroc_ge_0.80", {}).get("pass") else "FAIL",
     "python validation/phase1_seeded_anomaly.py"],
    ["Phase 1", "", "", "A2 beats naive RMSD-from-mean baseline", "> baseline",
     f"0.831 vs 1.000", "FAIL", ""],
    ["Phase 1", "", "", "A3 null control near chance", "0.40-0.60",
     a1.get("A3_null_control_near_chance", {}).get("value"), "PASS", ""],
    ["Phase 2", "Temporal scrambling (KINETIC anomaly)", "Block permutation junction list; ensemble unchanged",
     "B1 surprise AUROC >= 0.75 (near mode)", "0.75",
     a2.get("B1_surprise_near_auroc_ge_0.75", {}).get("value"), "FAIL",
     "python validation/phase2_temporal_scramble.py"],
    ["Phase 2", "", "", "B2 beats naive frame-to-frame jump detector", "> baseline",
     f"{a2.get('B2_beats_naive_frame_jump',{}).get('surprise')} vs {a2.get('B2_beats_naive_frame_jump',{}).get('frame_to_frame_rmsd')}",
     "PASS", ""],
    ["Phase 2", "", "", "B3 all time-blind baselines at chance", "0.40-0.60",
     "rarity 0.611 (misclassified)", "FAIL", ""],
    ["Phase 2", "", "", "B4 identity-null control at chance", "0.40-0.60",
     a2.get("B4_identity_null_at_chance", {}).get("surprise"), "PASS", ""],
    ["Phase 3", "Scale-up on ATLAS long trajectories", "Same as Phase 2, many proteins",
     "C1 median surprise AUROC >= 0.75", "0.75",
     a3.get("C1_median_surprise_ge_0.75", {}).get("value"), "N/A n=1",
     "python validation/phase3_scale.py --data_root ~/atlas"],
    ["Phase 3", "", "", "C2 beats best time-blind on >=80% of proteins", "0.80",
     a3.get("C2_beats_best_blind_on_80pct", {}).get("win_fraction"), "N/A n=1", ""],
    ["Phase 3", "", "", "C3 time-blind medians in band", "0.40-0.60", "not met", "N/A n=1", ""],
    ["Phase 3", "", "", "C4 Wilcoxon p < 0.05", "0.05",
     a3.get("C4_wilcoxon_p_lt_0.05", {}).get("p_value"), "N/A n=1", ""],
    ["Phase 3", "", "", "C5 identity-null control at chance (gates all others)", "0.40-0.60",
     "0.492-0.527", "PASS", ""],
]
end = rows(ws, design, wrap=(2, 3, 4, 8))
flag(ws, 7, 4, end, good_vals=("PASS",), bad_vals=("FAIL", "RETRACTED"))
for r in range(4, end):
    if ws.cell(row=r, column=7).value in ("PARTIAL", "N/A n=1"):
        ws.cell(row=r, column=7).fill = WARN
        ws.cell(row=r, column=7).font = B


# =====================================================================
# 2. DEFECTS  +  4. EDGE CASES
# =====================================================================
ws = sheet("02_Defects", "All defects and observations found",
           "Nine code defects (D-01..D-09), three observations (O-09..O-11), six errata (E-1..E-6), "
           "four phase findings (P1-*, P2-*). Full detail in VALIDATION_REPORT.md.",
           ["ID", "Severity", "Component", "Defect", "Measured impact", "Status"],
           [7, 10, 26, 42, 56, 16])
DEF = [
 ["D-01","Critical","scoring/anomaly_v2.py","Cluster labels never remapped to MSM active set",
  "28-50% of frames read the wrong stationary probability; rarity AUROC 0.45-0.80; 0% overlap in top-10% flagged frames","Fixed"],
 ["D-02","High","scoring/anomaly_v2.py","Rank normalisation broke ties by frame order",
  "Tie groups of 21-410 frames spread across 20-42 score points by index alone","Fixed"],
 ["D-03","High","scoring/anomaly_v2.py","Disconnected-state transitions scored 0 (least anomalous)",
  "Corrected-surprise AUROC 0.02 on 1CRN rare transitions (inverse ranking)","Fixed"],
 ["D-04","High","run_all_proteins.py","Per-residue scores reduced to pure RMSF",
  "Pearson r = 1.0000000000 vs RMSF pre-fix; 0.990-0.996 post-fix (see E-1)","Partially fixed"],
 ["D-05","High","run_all_proteins.py","Residue keys collided across chains",
  "8H0R: 182 residues collapsed to 91 entries; chain A overwritten by chain B","Fixed"],
 ["D-06","High","tools/export_for_asvs.py","ASVS export produced constant values",
  "anomaly_residue.json = 0.25 for all 182 residues x 100 frames","Fixed"],
 ["D-07","High","msm/validation.py","CK test compared mismatched state spaces",
  "Test invalid whenever any state was pruned - i.e. in every run performed","Fixed"],
 ["D-08","High","msm/select_lag_and_dim.py","VAMP-2 scorer numerically unstable on real features",
  "Returned 2.5e5 where a valid VAMP-2 score is bounded by the dimension (5)","Fixed"],
 ["D-09","Critical","run_all_proteins.py / batch_runner.py","Default smoothing erased isolated anomalies",
  "Scattered-mode AUROC 0.839 unsmoothed vs 0.550 at window=3","Fixed"],
 ["O-09","Info","run_all_proteins.py","Mean frame score ~50 is a normalisation artifact",
  "47.4-50.9 across all systems regardless of protein or dynamics","Documented"],
 ["O-10","Medium","deeptime KMeans","Clustering not reproducible across machines",
  "ARI 0.43-0.51 vs reference; tICA reproduces at r=0.96-0.9996","Open"],
 ["O-11","Critical","Sampling","100-frame trajectories cannot support MSM validity claims",
  "CK intervals +/-0.3-0.5; ITS unconverged; pi CI 1.1-1.6x the value","Open"],
 ["E-1","High","run_all_proteins.py","ERRATUM: residue scores still ~99% collinear with RMSF after D-04 fix",
  "Post-fix Pearson r = 0.996 / 0.990 / 0.992","Open"],
 ["E-2","Critical","validation/layer3_external.py","ERRATUM: B-factor correlation fully explained by RMSF",
  "Partial Spearman(B, score | RMSF) = -0.019 / -0.181","Open"],
 ["E-3","Critical","validation/layer3_external.py","ERRATUM: rare-STATE baseline comparison is circular",
  "Without rarity: loses to LOF on 8H0R (0.83 vs 0.93) and IsoForest on 1CRN (0.76 vs 0.85)","Open"],
 ["E-4","Critical","validation/layer3_external.py","ERRATUM: rare-TRANSITION headline also circular",
  "Spearman(surprise, -transition count) = 0.69 / 0.87","Open"],
 ["E-5","Medium","validation/layer2_msm_validity.py","ERRATUM: VAMP-2 saturation mechanism mis-stated",
  "20 test frames leave only 5 lagged pairs at lag 15 - degenerate, not merely overfit","Documented"],
 ["E-6","Info","VALIDATION_REPORT.md","ERRATUM: edge-case count stated as 21; actual 25","-","Documented"],
 ["P1-1","Critical","Method scope","Pipeline loses to a trivial baseline on seeded structural anomalies",
  "rmsd_from_mean 1.000 vs pipeline_fused 0.831; criterion A2 FAILED","Open"],
 ["P1-2","High","scoring/anomaly_v2.py","Transition surprise is BELOW chance on sustained rare states",
  "Burst-mode surprise_only AUROC 0.351","Open"],
 ["P2-1","High","Method scope","Kinetic channel detects geometrically invisible transitions (B1 marginally missed)",
  "surprise 0.741 +/- 0.113 vs frame_to_frame_rmsd 0.500 and time-blind 0.471-0.546","Partially fixed"],
 ["P2-3","Medium","Pre-registration","rarity mis-classified as time-blind in the B3 criterion",
  "rarity 0.611 broke the 0.40-0.60 band; all genuinely time-blind detectors were in band","Open"],
]
end = rows(ws, DEF, wrap=(4, 5))
for r in range(4, end):
    c = ws.cell(row=r, column=6); c.font = B
    c.fill = GOOD if c.value == "Fixed" else (BAD if c.value == "Open" else WARN)
    sv = ws.cell(row=r, column=2)
    if sv.value == "Critical": sv.fill = BAD; sv.font = B
    elif sv.value == "High": sv.fill = WARN

eb = load("edge_results_BASELINE.json") or {}
ep = load("edge_results_POSTFIX.json") or {}
ws = sheet("04_L1_EdgeCases", "Layer 1 — edge cases against the PRODUCTION code path",
           "Sources: edge_results_BASELINE.json / _POSTFIX.json. 25 cases. Note the repo's own edge-case suite "
           "tests scoring/signals.py, a parallel implementation the batch pipeline never calls.",
           ["Case", "Baseline status", "Post-fix status", "Post-fix detail"], [42, 14, 14, 74])
ed = []
for k in sorted(set(list(eb) + list(ep))):
    ed.append([k, (eb.get(k) or {}).get("status", "-"), (ep.get(k) or {}).get("status", "-"),
               str((ep.get(k) or {}).get("detail", ""))[:300]])
end = rows(ws, ed, wrap=(4,))
for r in range(4, end):
    for col in (2, 3):
        c = ws.cell(row=r, column=col)
        if c.value == "ok": c.fill = GOOD
        elif c.value == "CRASH": c.fill = WARN
ws.cell(row=end + 2, column=1, value="The two CRASH rows are constant / near-zero-variance features: deeptime raises "
        "ZeroRankError and the pipeline marks the protein failed with a logged cause. That is correct fail-fast "
        "behaviour, not a defect.").font = SM

# =====================================================================
# 3. L1 NUMERICAL
# =====================================================================
num = load("numerical_results_BASELINE.json") or {}
ws = sheet("03_L1_Numerical", "Layer 1 — numerical validation (baseline code)",
           "Source: numerical_results_BASELINE.json. Quantifies the state-indexing defect and the tie artifact.",
           ["Protein", "Frames", "KMeans labels", "MSM states", "Dropped labels",
            "Frames in dropped states", "Survivors misindexed", "Rarity Spearman (buggy vs correct)",
            "Rarity top-10% overlap", "Fused top-10% overlap", "Mean fused score"],
           [10, 9, 12, 11, 14, 15, 15, 18, 14, 14, 12])
data = []
for pid in ("1VII", "8H0R", "1CRN"):
    a = num.get(pid, {}).get("active_set", {})
    if not a:
        continue
    data.append([pid, a.get("n_frames"), a.get("n_clusters_labels"), a.get("n_msm_states"),
                 str(a.get("dropped_original_labels")), a.get("frames_in_dropped_states"),
                 a.get("frames_wrong_pi_row (survivors misindexed)"),
                 round(a.get("rarity_spearman", 0), 3), a.get("rarity_top10_overlap"),
                 a.get("fused_top10_overlap"), round(num[pid].get("mean_fused_score", 0), 2)])
end = rows(ws, data, num_fmt="0.000")

ws.cell(row=end + 2, column=1, value="TIE ARTIFACT (rank normalisation broke ties by frame order)").font = B
tie_hdr = ["Protein", "Largest tie group (rarity)", "Score spread", "Fair spread",
           "Largest tie group (surprise)", "Score spread", "Fair spread"]
for i, h in enumerate(tie_hdr, 1):
    c = ws.cell(row=end + 3, column=i, value=h); c.font = B; c.fill = SUB; c.border = BORDER
td = []
for pid in ("1VII", "8H0R", "1CRN"):
    t = num.get(pid, {}).get("ties", {})
    if not t:
        continue
    td.append([pid, t.get("rarity_largest_tie_group_size"),
               round(t.get("rarity_score_spread_within_tie_group", 0), 1),
               t.get("rarity_fair_spread_within_tie_group"),
               t.get("surprise_largest_tie_group_size"),
               round(t.get("surprise_score_spread_within_tie_group", 0), 1),
               t.get("surprise_fair_spread_within_tie_group")])
end2 = rows(ws, td, start=end + 4)

ws.cell(row=end2 + 2, column=1, value="UNIT CHECKS (analytic 3-state chain, hand-computed)").font = B
uc = num.get("unit_checks_tiny_msm", {})
r = end2 + 3
for k in ("rarity_state0_expected", "rarity_state0_actual", "surprise_0to1_expected",
          "surprise_0to1_actual", "surprise_2to2_expected", "surprise_2to2_actual",
          "rarity_math_ok", "surprise_math_ok"):
    if k in uc:
        ws.cell(row=r, column=1, value=k).font = N
        ws.cell(row=r, column=2, value=uc[k]).font = N
        r += 1

# =====================================================================
# 5. ACCURACY
# =====================================================================
rb = load("rare_case_results_BASELINE.json") or {}
rp = load("rare_case_results_POSTFIX.json") or {}
ws = sheet("05_Accuracy", "Detection accuracy — baseline vs post-fix",
           "Sources: rare_case_results_BASELINE.json / _POSTFIX.json. "
           "NOTE: ground truth is MSM-derived; see sheet 08 for the circularity finding.",
           ["Protein", "Frames", "Task", "Signal", "AUROC baseline", "AUROC post-fix",
            "Delta", "AUPRC baseline", "AUPRC post-fix", "n positives"],
           [9, 8, 16, 26, 14, 14, 10, 14, 14, 11])
acc = []
keymap = [("rare_state | fused (as-shipped)", "Rare state", "Fused score"),
          ("rare_state | rarity channel (as-shipped)", "Rare state", "Rarity channel"),
          ("rare_trans | fused (as-shipped)", "Rare transition", "Fused score"),
          ("rare_trans | surprise channel (as-shipped)", "Rare transition", "Surprise channel")]
for pid in rb:
    for key, task, sig in keymap:
        b = rb[pid]["detection"].get(key, {})
        p = rp.get(pid, {}).get("detection", {}).get(key, {})
        if not b:
            continue
        acc.append([pid, rb[pid].get("n_frames"), task, sig, b.get("auroc"), p.get("auroc"),
                    round((p.get("auroc") or 0) - (b.get("auroc") or 0), 3),
                    b.get("auprc"), p.get("auprc"), b.get("n_pos")])
end = rows(ws, acc, num_fmt="0.000")
ws.cell(row=end + 2, column=1,
        value="Rarity-channel rows reaching 1.000 are circular by construction (labels and channel both derive from pi). "
              "Reported as a consistency check, not an accuracy claim.").font = SM

# =====================================================================
# 6. L2
# =====================================================================
l2 = load("layer2_results.json") or {}
ws = sheet("06_L2_MSM_Validity", "Layer 2 — MSM statistical validity",
           "Source: layer2_results.json. Executed on our own systems (previously only mock data).",
           ["System", "Frames", "CK: fraction within 95% CI", "CK: max deviation",
            "ITS: rel change at last lag step", "Bootstrap median CI width",
            "CI width / pi", "VAMP-2 best score", "VAMP-2 saturated", "Verdict"],
           [9, 8, 16, 13, 17, 16, 11, 13, 12, 24])
l2d = []
frames = {"1VII": 100, "8H0R": 100, "1UBQ": 100, "1CRN": 1001}
for pid, v in l2.items():
    ck = v.get("chapman_kolmogorov", {}); its = v.get("implied_timescales", {})
    bs = v.get("bootstrap_pi", {}); vp = v.get("vamp2_grid", {})
    best = vp.get("best", {}) if "best" in vp else {}
    ciw = bs.get("median_ci_width_over_pi")
    itsc = its.get("slowest_its_rel_change_last_step")
    verdict = ("Publication-grade" if (ciw or 9) < 0.5 and (itsc or 9) < 0.1
               else "Insufficient sampling")
    l2d.append([pid, frames.get(pid), ck.get("fraction_within_95CI"),
                max(ck.get("max_abs_deviation_per_state", [0])), itsc,
                bs.get("median_ci_width"), ciw, best.get("score"),
                "YES" if best.get("saturated_at_dim (overfit flag)") else "n/a", verdict])
end = rows(ws, l2d, num_fmt="0.000")
for r in range(4, end):
    c = ws.cell(row=r, column=10)
    c.fill = BAD if c.value == "Insufficient sampling" else GOOD
    c.font = B
ws.cell(row=end + 2, column=1, value="ACCEPTANCE CRITERIA").font = B
for i, t in enumerate([
        "Chapman-Kolmogorov: fraction within 95% CI >= 0.90 AND intervals narrow enough to be informative",
        "Implied timescales: relative change over final lag step < 0.10 (plateau reached)",
        "Bootstrap pi: median CI width / pi < 0.50",
        "VAMP-2: optimum strictly below the dimension bound (no saturation)"]):
    ws.cell(row=end + 3 + i, column=1, value=t).font = SM

# =====================================================================
# 7. L3
# =====================================================================
l3 = load("layer3_results.json") or {}
ws = sheet("07_L3_External", "Layer 3 — external validation",
           "Source: layer3_results.json. SEE SHEET 08: the B-factor claim was RETRACTED and the "
           "baseline comparison shown to be circular.",
           ["Block", "System", "Method / detector", "Metric", "Value", "Note"],
           [14, 10, 32, 26, 12, 52])
l3d = []
for pid, v in (l3.get("part_A_bfactor") or {}).items():
    if "skipped" in v:
        l3d.append(["B-factor", pid, v.get("method", ""), "-", None, v["skipped"]])
        continue
    l3d.append(["B-factor", pid, v.get("method", ""), "Spearman vs dynamic score",
                v.get("spearman_B_vs_dynamic_score"),
                "RETRACTED: partial correlation controlling for RMSF is ~0 (see sheet 08)"])
    l3d.append(["B-factor", pid, v.get("method", ""), "Spearman vs trajectory RMSF",
                v.get("spearman_B_vs_traj_RMSF"), "This is the signal actually being validated"])
for pid, v in (l3.get("part_B_baselines") or {}).items():
    for meth, m in v.items():
        if not isinstance(m, dict):
            continue
        l3d.append(["Baseline", pid, meth, "rare-state AUROC", m.get("rare_state_AUROC"),
                    "CIRCULAR: labels derive from pi, rarity channel IS 1-pi (sheet 08)"])
        l3d.append(["Baseline", pid, meth, "rare-transition AUROC", m.get("rare_trans_AUROC"), ""])
end = rows(ws, l3d, wrap=(6,), num_fmt="0.000")

# =====================================================================
# 8. REVALIDATION
# =====================================================================
rv = load("revalidation_results.json") or {}
ws = sheet("08_Revalidation", "Adversarial re-validation — our own claims, re-tested to falsify",
           "Source: revalidation_results.json. THIS SHEET RETRACTS THREE HEADLINE CLAIMS. Read before citing sheets 05 or 07.",
           ["Check", "What was tested", "Result", "Verdict"], [10, 40, 40, 70])
rvd = []
labels = {
    "R1_residue_score_vs_rmsf": "Did the residue-score fix decouple scores from RMSF?",
    "R2_chain_counts": "Multi-chain residue count after the key fix",
    "R3_partial_bfactor": "Does the dynamic score add anything beyond RMSF vs B-factors?",
    "R4_state_label_circularity": "Is the rare-STATE baseline comparison circular?",
    "R5_transition_label_circularity": "Is the rare-TRANSITION headline circular?",
    "R6_vamp2_test_pairs": "VAMP-2 saturation - real mechanism?",
    "R7_edge_case_count": "Edge-case count consistency",
    "R8_core_finding_recheck": "Independent re-derivation of the core defect (D-01)",
}
for k, v in rv.items():
    if not isinstance(v, dict):
        continue
    detail = {kk: vv for kk, vv in v.items() if kk != "_verdict"}
    rvd.append([k.split("_")[0], labels.get(k, k),
                json.dumps(detail)[:600], v.get("_verdict", "")])
end = rows(ws, rvd, wrap=(2, 3, 4))

# =====================================================================
# 9 / 10. PHASE 1 & 2
# =====================================================================
def phase_sheet(name, title, sub, res, mode_key="modes"):
    ws = sheet(name, title, sub,
               ["Mode", "Method", "AUROC mean", "AUROC std", "AUPRC mean", "Time-blind?", "Per-replicate AUROC"],
               [16, 26, 12, 11, 12, 12, 40])
    TB = ["rmsd_from_mean", "feature_zscore", "isolation_forest",
          "local_outlier_factor", "density_only", "rarity_only"]
    d = []
    for mode, mv in (res.get(mode_key) or {}).items():
        for meth, m in sorted((mv.get("methods") or {}).items(),
                              key=lambda kv: -kv[1]["auroc_mean"]):
            d.append([mode, meth, m.get("auroc_mean"), m.get("auroc_std"),
                      m.get("auprc_mean"), "yes" if meth in TB else "",
                      str(m.get("auroc_per_replicate", ""))[:90]])
    return ws, rows(ws, d, num_fmt="0.000")


ws, end = phase_sheet("09_Phase1_Seeded",
    "Phase 1 — seeded THERMODYNAMIC anomaly (450 K frames in 300 K background)",
    "Source: phase1_results.json. Ground truth = splice list, fully independent of the MSM. 5 replicates.",
    p1)
ws.cell(row=end + 2, column=1, value="KEY FINDING: the pipeline LOST to a trivial baseline. "
        "rmsd_from_mean AUROC 1.000 vs pipeline_fused 0.831. Criterion A2 FAILED. "
        "450 K frames are globally expanded, which is the case geometry owns.").font = B
ws.cell(row=end + 3, column=1, value="SECOND FINDING: surprise scored 0.351 (BELOW chance) on sustained rare states - "
        "within a contiguous rare block, transitions are self-consistent so P is high and surprise LOW. "
        "The channel detects state CHANGES, not residence in a rare state.").font = SM
ws.cell(row=end + 4, column=1, value="THIRD FINDING (defect D-09): default smoothing erased isolated anomalies - "
        "scattered-mode AUROC 0.839 unsmoothed vs 0.550 at window=3. Now fixed (smoothing opt-in).").font = SM

ws, end = phase_sheet("10_Phase2_Scramble",
    "Phase 2 — temporal scrambling (the decisive KINETIC test)",
    "Source: phase2_results.json. Blocks permuted so every frame is used exactly once: the static ensemble is "
    "IDENTICAL to the original and only frame ORDER differs. 12 replicates.",
    p2)
ws.cell(row=end + 2, column=1, value="KEY RESULT: in near mode (junctions geometrically smooth by construction), "
        "transition surprise 0.741 is the only method above chance. The naive frame-to-frame jump detector sits at "
        "exactly 0.500 and every time-blind detector at 0.471-0.546.").font = B
ws.cell(row=end + 3, column=1, value="Identity-null control: surprise 0.501 - confirms the construction generates no "
        "artificial signal. B2 and B4 PASSED; B1 missed the 0.75 bar by 0.009 (inside the +/-0.113 spread); "
        "B3 failed only because rarity was mis-classified as time-blind (pi is estimated from transition counts).").font = SM

# =====================================================================
# 11. PHASE 3
# =====================================================================
ws = sheet("11_Phase3_ATLAS", "Phase 3 — scale-up pilot on ATLAS long trajectories",
           "Source: phase3_results.json. ATLAS 1g2r_A, 3 replicates x 1000 frames. "
           "n=1 INDEPENDENT PROTEIN: C1/C3/C4 are not interpretable at this sample size.",
           ["Trajectory", "Frames", "Method", "AUROC (near mode)", "AUROC (identity null)",
            "Time-blind?", "Drift first-vs-last decile (nm)"], [24, 8, 24, 15, 16, 11, 18])
TB = ["rmsd_from_mean", "feature_zscore", "isolation_forest",
      "local_outlier_factor", "density_only", "rarity_only"]
p3d = []
for sid, v in (p3.get("per_system") or {}).items():
    for meth in sorted(v.get("auroc", {}), key=lambda m: -v["auroc"][m]):
        p3d.append([sid, 1000, meth, v["auroc"].get(meth),
                    (v.get("auroc_null") or {}).get(meth),
                    "yes" if meth in TB else "",
                    v.get("drift_rmsd_first_vs_last_decile")])
end = rows(ws, p3d, num_fmt="0.000")
ws.cell(row=end + 2, column=1, value="NULL CONTROL PASSED (all methods 0.492-0.527). The construction is valid on long, "
        "drifting trajectories - junctions carry no positional bias. This was the open question after Phase 2.").font = B
ws.cell(row=end + 3, column=1, value="Per-replicate spread is very large for geometric detectors on the SAME protein "
        "(isolation_forest 0.395 / 0.525 / 0.865). At n=1 protein the medians are arithmetic on a single number. "
        "~20 proteins are needed before C1/C3/C4 mean anything.").font = SM

# =====================================================================
# 12. D09
# =====================================================================
d09 = load("d09_impact.json") or {}
ws = sheet("12_D09_Impact", "D-09 — effect of the smoothing fix on earlier accuracy results",
           "Source: d09_impact.json. Smoothing is now opt-in (default OFF).",
           ["Protein", "Rare-state w=3", "Rare-state w=1", "Delta",
            "Rare-transition w=3", "Rare-transition w=1", "Delta"], [10, 14, 14, 10, 16, 16, 10])
dd = []
for pid, v in d09.items():
    dd.append([pid, v.get("state_w3"), v.get("state_w1"),
               round((v.get("state_w1") or 0) - (v.get("state_w3") or 0), 3),
               v.get("trans_w3"), v.get("trans_w1"),
               round((v.get("trans_w1") or 0) - (v.get("trans_w3") or 0), 3)])
end = rows(ws, dd, num_fmt="0.000")
ws.cell(row=end + 2, column=1, value="Turning smoothing off barely changes these numbers (mean delta -0.010 / +0.017) because the "
        "MSM-derived labels mark SUSTAINED rare states. The 0.84 -> 0.55 collapse appears only for ISOLATED events "
        "(Phase 1, sheet 09) - which the old ground truth structurally could not contain.").font = B

# =====================================================================
# 13. LIMITATIONS
# =====================================================================
ws = sheet("13_Limitations", "Known limitations and open questions",
           "Stated explicitly so a reviewer does not have to find them.",
           ["#", "Limitation", "Consequence", "Status"], [5, 52, 62, 20])
lim = [
    [1, "Rare-case ground truth in sheet 05 is MSM-derived (bottom-quartile-pi states, low-count transitions)",
     "Those accuracy figures measure internal consistency, not biological validity. The baseline comparison in sheet 07 is circular: labels come from pi and the rarity channel IS 1-pi.", "Open (OI-12)"],
    [2, "Per-residue scores remain ~99% collinear with RMSF (r = 0.990-0.996) after the D-04 fix",
     "The B-factor correlation in sheet 07 validates RMSF, not the anomaly signal (partial rho ~0). Residue-level claims are not currently supportable.", "Open (OI-13)"],
    [3, "Trajectories used for Layers 1-2 are 100-frame implicit-solvent toy runs (1001 for 1CRN)",
     "CK intervals +/-0.3-0.5 (no statistical power); implied timescales unconverged; pi uncertain to +/-100%; VAMP-2 saturates on a degenerate 5-pair test split.", "Open (OI-1)"],
    [4, "Phase 3 has n=1 independent protein",
     "C1, C3 and C4 are arithmetic on a single number. Only C5 (null control) is interpretable, and it passed.", "In progress"],
    [5, "Phase 1 showed the method loses to RMSD-from-mean on structural anomalies (0.831 vs 1.000)",
     "Scope must be limited to KINETIC anomalies. For structural outliers a simpler method is better and the paper should say so.", "Documented (P1-1)"],
    [6, "Transition surprise scores BELOW chance (0.351) on sustained rare states",
     "The channel detects state CHANGES, not residence in a rare state. Conceptual mismatch with 'rare-case detection' as originally framed.", "Documented (P1-2)"],
    [7, "deeptime KMeans is not reproducible across machines despite a fixed seed (ARI 0.43-0.51)",
     "State assignments are platform-specific; only score-level and AUROC-level claims are portable.", "Open (O-10)"],
    [8, "Phase 2 B3 criterion mis-classified rarity as time-blind",
     "pi is estimated from the transition count matrix, so rarity inherits temporal information. Recorded as a pre-registration flaw; a corrected replication is still owed.", "Open (P2-3)"],
    [9, "Mean frame score is ~50 for every protein",
     "Rank normalisation forces a uniform distribution per channel. Must never be reported as a quality or health metric.", "Documented (O-09)"],
    [10, "ATLAS data is CC-BY-NC",
     "Non-commercial only. Relevant if the visualisation software is ever commercialised.", "Noted"],
]
end = rows(ws, lim, wrap=(2, 3))
for r in range(4, end):
    c = ws.cell(row=r, column=4)
    c.font = B
    c.fill = BAD if str(c.value).startswith("Open") else WARN

wb.save("/home/claude/report/SciML-MD_Validation_Dataset.xlsx")
print("saved")
